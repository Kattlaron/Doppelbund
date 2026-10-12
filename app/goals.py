from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


GOALS_FILE = Path(__file__).resolve().parents[1] / "assets" / "cards" / "cards.xlsx"


@dataclass(frozen=True)
class Goal:
    id: str
    condition: str
    reward: int


def load_goals() -> list[Goal]:
    workbook = load_workbook(GOALS_FILE, read_only=True, data_only=True)
    sheet = workbook["goals"]
    goals = [
        Goal(str(name), str(condition), int(reward))
        for name, condition, reward in sheet.iter_rows(min_row=2, values_only=True)
        if name and condition and reward is not None
    ]
    workbook.close()
    return goals

