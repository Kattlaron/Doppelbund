from __future__ import annotations

from dataclasses import dataclass
from functools import lru_cache
from pathlib import Path

from openpyxl import load_workbook


EVENTS_FILE = Path(__file__).resolve().parents[1] / "assets" / "cards" / "cards.xlsx"

IMPLEMENTED_EVENT_IDS = frozenset({
    "storm",
    "harvest",
    "darkness",
    "rebellion",
    "cycle",
    "prophecy",
    "silence",
    "wandering_throne",
    "revelation",
    "mystery",
    "balance",
})


@dataclass(frozen=True)
class GameEvent:
    id: str
    name: str
    effect: str


@lru_cache(maxsize=1)
def load_supported_events() -> tuple[GameEvent, ...]:
    workbook = load_workbook(EVENTS_FILE, read_only=True, data_only=True)
    try:
        sheet = workbook["event"]
        events = tuple(
            GameEvent(str(event_id).strip(), str(name).strip(), str(effect).strip())
            for event_id, name, effect in sheet.iter_rows(min_row=2, values_only=True)
            if event_id and name and effect and str(event_id).strip() in IMPLEMENTED_EVENT_IDS
        )
    finally:
        workbook.close()

    found_ids = {event.id for event in events}
    missing_ids = IMPLEMENTED_EVENT_IDS - found_ids
    if missing_ids:
        raise ValueError(
            "Im Tabellenblatt 'event' fehlen implementierte Event-IDs: "
            + ", ".join(sorted(missing_ids))
        )
    return events
